#!/usr/bin/env python3
"""Extract text from PDF, DOCX, and PPTX attachments.

Emits the same JSONL shape as ocr_vision.py so the verifier and write-back stages
consume images and documents through one path.

PDFs are the interesting case: most carry a real text layer, which is *exact* rather
than recognized -- no OCR error to verify. But a scanned PDF has no text layer at all,
and silently emitting an empty string for it would look like a successful extraction of
a blank document. So each page is checked: text layer if present, rendered-and-OCR'd if
not. The `method` field records which happened, per page, because a page recovered by
OCR carries recognition risk that a text-layer page does not.

    python3 extract_docs.py --dir /path/to/attachments --out docs.jsonl
"""
from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path

SCANNED_PAGE_THRESHOLD = 50   # chars below which a page is treated as image-only
RENDER_DPI = 200              # rasterize scanned pages at this DPI before OCR


def extract_pdf(path: Path) -> dict:
    """Text layer where present; Vision OCR on rendered pages where absent."""
    import fitz  # pymupdf
    try:
        from capture.pages.ocr_vision import recognize
    except ImportError:                                   # run as a script from reference/
        import sys
        sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
        from capture.pages.ocr_vision import recognize

    rec = {"path": str(path), "kind": "pdf", "ok": False}
    try:
        doc = fitz.open(path)
    except Exception as exc:  # noqa: BLE001
        return {**rec, "error": f"{type(exc).__name__}: {exc}"}

    parts, methods, confs = [], [], []
    try:
        for page in doc:
            text = page.get_text().strip()
            if len(text) >= SCANNED_PAGE_THRESHOLD:
                parts.append(text)
                methods.append("text_layer")
                confs.append(1.0)          # a text layer is exact, not recognized
                continue
            # No usable text layer -- this page is an image. Render and OCR it.
            with tempfile.NamedTemporaryFile(suffix=".png", delete=True) as tmp:
                page.get_pixmap(dpi=RENDER_DPI).save(tmp.name)
                ocr = recognize(Path(tmp.name))
            if ocr.get("ok") and ocr.get("chars", 0) > 0:
                parts.append(ocr["text"])
                methods.append("ocr")
                confs.append(ocr.get("mean_conf", 0.0))
            else:
                parts.append("")
                methods.append("empty")
    finally:
        doc.close()

    text = "\n\n".join(p for p in parts if p)
    return {
        **rec,
        "ok": True,
        "pages": len(parts),
        "methods": methods,
        "method": "mixed" if len(set(methods)) > 1 else (methods[0] if methods else "empty"),
        "mean_conf": round(sum(confs) / len(confs), 4) if confs else 0.0,
        "chars": len(text),
        "text": text,
    }


def extract_docx(path: Path) -> dict:
    """Paragraphs plus table cells -- tables carry real content in meeting notes."""
    import docx

    rec = {"path": str(path), "kind": "docx", "ok": False, "method": "text_layer"}
    try:
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        return {**rec, "ok": True, "mean_conf": 1.0, "chars": len(text), "text": text}
    except Exception as exc:  # noqa: BLE001
        return {**rec, "error": f"{type(exc).__name__}: {exc}"}


def extract_pptx(path: Path) -> dict:
    """Slide text frames plus speaker notes, labelled by slide number."""
    from pptx import Presentation

    rec = {"path": str(path), "kind": "pptx", "ok": False, "method": "text_layer"}
    try:
        prs = Presentation(str(path))
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            body = [sh.text_frame.text.strip() for sh in slide.shapes
                    if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
            if slide.has_notes_slide:
                note = slide.notes_slide.notes_text_frame.text.strip()
                if note:
                    body.append(f"[notes] {note}")
            if body:
                parts.append(f"--- slide {i} ---\n" + "\n".join(body))
        text = "\n\n".join(parts)
        return {**rec, "ok": True, "slides": len(prs.slides), "mean_conf": 1.0,
                "chars": len(text), "text": text}
    except Exception as exc:  # noqa: BLE001
        return {**rec, "error": f"{type(exc).__name__}: {exc}"}


HANDLERS = {".pdf": extract_pdf, ".docx": extract_docx, ".pptx": extract_pptx}


def main() -> int:
    ap = argparse.ArgumentParser(description="Extract text from PDF/DOCX/PPTX -> JSONL")
    ap.add_argument("--dir", required=True, help="directory to walk")
    ap.add_argument("--out", required=True)
    ap.add_argument("--progress-every", type=int, default=10)
    args = ap.parse_args()

    paths = sorted(p for p in Path(args.dir).rglob("*")
                   if p.is_file() and p.suffix.lower() in HANDLERS)
    if not paths:
        print("no PDF/DOCX/PPTX found", file=sys.stderr)
        return 1
    print(f"extracting {len(paths)} documents", file=sys.stderr)

    with open(args.out, "w", encoding="utf-8") as fh:
        for i, p in enumerate(paths, 1):
            rec = HANDLERS[p.suffix.lower()](p)
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
            fh.flush()
            if args.progress_every and i % args.progress_every == 0:
                print(f"  {i}/{len(paths)}", file=sys.stderr)
    print(f"done: {len(paths)} documents", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
